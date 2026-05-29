"""
mmt_builder.py
==============
Word (.docx) generation engine for the Medical Missionary Training course.

Every lesson is authored as a plain Python dictionary (see lessons/lesson_1_1.py
for the schema) and rendered into a styled, self-contained Word document by
`build_lesson()`.  The goal is a consistent, professional, print-ready study
guide: cover/title block, a "Lesson at a Glance" panel, a memory verse, richly
formatted body content (headings, Scripture blocks, Spirit-of-Prophecy quote
blocks, lists, tables, callouts, figures), a summary, review questions, a
practical assignment, and a full reference list.

No external converter is required -- output is generated directly with
python-docx, so it opens natively in Microsoft Word, Google Docs and LibreOffice.
"""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ───────────────────────────── THEME ──────────────────────────────────────
BODY_FONT      = "Georgia"
HEAD_FONT      = "Georgia"

GREEN_DARK     = "14532D"   # H1 / title       (deep forest green)
GREEN          = "256D43"   # H2
GREEN_SOFT     = "2E7D52"   # accents / quote border
GOLD_DARK      = "8A6D00"   # H3 / scripture reference
GOLD           = "B8860B"   # scripture border
INK            = "222222"   # body text
GREY           = "5B5B5B"   # captions / notes
LIGHT_GREEN_BG = "EAF3EC"   # callout background
LIGHT_GOLD_BG  = "FBF6E7"   # scripture background
LIGHT_QUOTE_BG = "F1F6F2"   # SOP quote background
ROW_ALT_BG     = "F2F7F3"   # table alternate row
HEADER_BG      = "256D43"   # table header fill


# ──────────────────────── low-level XML helpers ───────────────────────────
def _shade(element_pr, color_hex):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    element_pr.append(shd)


def shade_paragraph(paragraph, color_hex):
    _shade(paragraph._p.get_or_add_pPr(), color_hex)


def shade_cell(cell, color_hex):
    _shade(cell._tc.get_or_add_tcPr(), color_hex)


def left_border(paragraph, color_hex, sz="26", space="14"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), sz)
    left.set(qn('w:space'), space)
    left.set(qn('w:color'), color_hex)
    pBdr.append(left)
    pPr.append(pBdr)


def bottom_border(paragraph, color_hex, sz="8"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def table_borders(table, color_hex, sz="6"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), sz)
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), color_hex)
        borders.append(e)
    tblPr.append(borders)


def cell_margins(cell, top=120, start=160, bottom=120, end=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = OxmlElement(f'w:{tag}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_page_number_field(paragraph):
    """Insert a live PAGE field so Word shows real page numbers."""
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin); run._r.append(instr); run._r.append(end)
    return run


# ───────────────────────── inline rich-text runs ──────────────────────────
_INLINE_RE = re.compile(r'(\*\*.+?\*\*|\*.+?\*)')


def add_runs(paragraph, text, italic=False, bold=False, color=None, size=None, font=BODY_FONT):
    """Add text to a paragraph, honouring **bold** and *italic* inline markers."""
    for tok in _INLINE_RE.split(text):
        if not tok:
            continue
        b, i, inner = bold, italic, tok
        if tok.startswith('**') and tok.endswith('**'):
            b, inner = True, tok[2:-2]
        elif tok.startswith('*') and tok.endswith('*'):
            i, inner = True, tok[1:-1]
        run = paragraph.add_run(inner)
        run.bold, run.italic = b, i
        run.font.name = font
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)
    return paragraph


# ──────────────────────────── block renderers ─────────────────────────────
def _new_para(container, space_before=2, space_after=8, line=1.18):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    return p


def render_heading(container, text, level):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 2 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    color, size = (GREEN, 14) if level == 2 else (GOLD_DARK, 12)
    add_runs(p, text, bold=True, color=color, size=size, font=HEAD_FONT)
    if level == 2:
        bottom_border(p, GREEN_SOFT, sz="6")


def render_paragraph(container, text):
    p = _new_para(container)
    add_runs(p, text, color=INK, size=11)


def render_list(container, items, ordered=False):
    for idx, item in enumerate(items, 1):
        p = container.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.32)
        pf.first_line_indent = Inches(-0.22)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.12
        marker = f"{idx}.  " if ordered else "•  "
        run = p.add_run(marker); run.bold = ordered; run.font.name = BODY_FONT
        run.font.color.rgb = RGBColor.from_string(GREEN_SOFT)
        add_runs(p, item, color=INK, size=11)


def render_scripture(container, ref, text):
    p = _new_para(container, space_before=6, space_after=2)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    left_border(p, GOLD)
    shade_paragraph(p, LIGHT_GOLD_BG)
    add_runs(p, f"“{text}”", italic=True, color=INK, size=11)
    pr = _new_para(container, space_before=0, space_after=8)
    pr.paragraph_format.left_indent = Inches(0.25)
    left_border(pr, GOLD)
    shade_paragraph(pr, LIGHT_GOLD_BG)
    add_runs(pr, f"— {ref} (KJV)", bold=True, color=GOLD_DARK, size=10)


def render_quote(container, text, attribution):
    p = _new_para(container, space_before=6, space_after=2)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    left_border(p, GREEN_SOFT)
    shade_paragraph(p, LIGHT_QUOTE_BG)
    add_runs(p, f"“{text}”", italic=True, color=INK, size=11)
    if attribution:
        pa = _new_para(container, space_before=0, space_after=8)
        pa.paragraph_format.left_indent = Inches(0.25)
        left_border(pa, GREEN_SOFT)
        shade_paragraph(pa, LIGHT_QUOTE_BG)
        add_runs(pa, f"— {attribution}", bold=True, color=GREEN, size=10)


def _strip_leading_empty(cell):
    first = cell.paragraphs[0]
    if not first.text.strip() and len(cell.paragraphs) > 1:
        first._p.getparent().remove(first._p)


def render_callout(container, title, blocks, bg=LIGHT_GREEN_BG, border=GREEN_SOFT, title_color=GREEN):
    table = container.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table, border, sz="6")
    cell = table.cell(0, 0)
    shade_cell(cell, bg)
    cell_margins(cell)
    if title:
        tp = cell.paragraphs[0]
        tp.paragraph_format.space_after = Pt(4)
        add_runs(tp, title, bold=True, color=title_color, size=11.5, font=HEAD_FONT)
    for blk in blocks:
        render_block(cell, blk)
    _strip_leading_empty(cell)
    container.add_paragraph().paragraph_format.space_after = Pt(2)


def render_figure(container, title, lines, bg=LIGHT_GOLD_BG, border=GOLD):
    """A centred 'figure' box used for diagrams / acronyms / mnemonics."""
    table = container.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table, border, sz="6")
    cell = table.cell(0, 0)
    shade_cell(cell, bg)
    cell_margins(cell, top=140, bottom=140)
    if title:
        tp = cell.paragraphs[0]
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tp.paragraph_format.space_after = Pt(6)
        add_runs(tp, title, bold=True, color=GOLD_DARK, size=11.5, font=HEAD_FONT)
    for line in lines:
        lp = cell.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.paragraph_format.space_after = Pt(2)
        add_runs(lp, line, color=INK, size=11)
    _strip_leading_empty(cell)
    container.add_paragraph().paragraph_format.space_after = Pt(2)


def render_table(container, header, rows):
    table = container.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table, GREEN_SOFT, sz="4")
    hdr = table.rows[0].cells
    for i, htext in enumerate(header):
        shade_cell(hdr[i], HEADER_BG)
        cell_margins(hdr[i], top=70, bottom=70)
        p = hdr[i].paragraphs[0]
        add_runs(p, htext, bold=True, color="FFFFFF", size=10.5, font=HEAD_FONT)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c_idx, val in enumerate(row):
            cell_margins(cells[c_idx], top=60, bottom=60)
            if r_idx % 2 == 1:
                shade_cell(cells[c_idx], ROW_ALT_BG)
            p = cells[c_idx].paragraphs[0]
            add_runs(p, str(val), color=INK, size=10)
    container.add_paragraph().paragraph_format.space_after = Pt(2)


def render_note(container, text):
    p = _new_para(container, space_before=4, space_after=8)
    p.paragraph_format.left_indent = Inches(0.1)
    add_runs(p, "Note: ", bold=True, color=GREY, size=10)
    add_runs(p, text, italic=True, color=GREY, size=10)


def render_hr(container):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    bottom_border(p, "CCCCCC", sz="6")


_DISPATCH = {
    "h2":        lambda c, b: render_heading(c, b[1], 2),
    "h3":        lambda c, b: render_heading(c, b[1], 3),
    "p":         lambda c, b: render_paragraph(c, b[1]),
    "ul":        lambda c, b: render_list(c, b[1], ordered=False),
    "ol":        lambda c, b: render_list(c, b[1], ordered=True),
    "scripture": lambda c, b: render_scripture(c, b[1], b[2]),
    "quote":     lambda c, b: render_quote(c, b[1], b[2]),
    "callout":   lambda c, b: render_callout(c, b[1], b[2]),
    "figure":    lambda c, b: render_figure(c, b[1], b[2]),
    "table":     lambda c, b: render_table(c, b[1], b[2]),
    "note":      lambda c, b: render_note(c, b[1]),
    "hr":        lambda c, b: render_hr(c),
}


def render_block(container, block):
    fn = _DISPATCH.get(block[0])
    if fn is None:
        raise ValueError(f"Unknown block type: {block[0]!r}")
    fn(container, block)


# ─────────────────────────── document chrome ──────────────────────────────
def _configure_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.line_spacing = 1.18
    pf.space_after = Pt(8)
    # ensure East-Asian fallback doesn't override the latin font
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(attr), BODY_FONT)


def _set_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)


def _header_footer(doc, meta):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    hdr = section.header.paragraphs[0]
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(hdr, "Medical Missionary Training", italic=True, color=GREY, size=8.5)

    ftr = section.footer.paragraphs[0]
    ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(ftr, f"Module {meta['module']} · Lesson {meta['number']}  ·  ", color=GREY, size=8.5)
    pnum = add_page_number_field(ftr)
    pnum.font.size = Pt(8.5); pnum.font.color.rgb = RGBColor.from_string(GREY); pnum.font.name = BODY_FONT


def _title_block(doc, meta):
    # Module eyebrow
    eb = doc.add_paragraph()
    eb.paragraph_format.space_after = Pt(2)
    add_runs(eb, f"MODULE {meta['module']}  —  {meta['module_title'].upper()}",
             bold=True, color=GOLD_DARK, size=10.5)
    # Lesson number
    ln = doc.add_paragraph()
    ln.paragraph_format.space_after = Pt(0)
    add_runs(ln, f"Lesson {meta['number']}", bold=True, color=GREEN_SOFT, size=12)
    # Title
    tp = doc.add_paragraph()
    tp.paragraph_format.space_before = Pt(2)
    tp.paragraph_format.space_after = Pt(2)
    add_runs(tp, meta['title'], bold=True, color=GREEN_DARK, size=22, font=HEAD_FONT)
    if meta.get('subtitle'):
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(4)
        add_runs(sp, meta['subtitle'], italic=True, color=GREY, size=12)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(8)
    bottom_border(rule, GOLD, sz="14")


def _glance_panel(doc, meta):
    blocks = []
    if meta.get('est_time'):
        blocks.append(("p", f"**Estimated study time:** {meta['est_time']}"))
    if meta.get('objectives'):
        blocks.append(("h3", "Learning Objectives"))
        blocks.append(("p", "By the end of this lesson, you will be able to:"))
        blocks.append(("ul", meta['objectives']))
    if meta.get('key_texts'):
        blocks.append(("h3", "Key Scripture Passages"))
        blocks.append(("ul", meta['key_texts']))
    if meta.get('key_sop'):
        blocks.append(("h3", "Key Spirit of Prophecy Sources"))
        blocks.append(("ul", meta['key_sop']))
    render_callout(doc, "\U0001F4D6  LESSON AT A GLANCE", blocks)


def _memory_verse(doc, meta):
    mv = meta.get('memory_verse')
    if not mv:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    add_runs(p, "MEMORY VERSE", bold=True, color=GOLD_DARK, size=10)
    render_scripture(doc, mv[0], mv[1])


def _section_title(doc, text, emoji=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    label = f"{emoji}  {text}" if emoji else text
    add_runs(p, label, bold=True, color=GREEN, size=15, font=HEAD_FONT)
    bottom_border(p, GREEN_SOFT, sz="8")


def _references(doc, meta):
    refs = meta.get('references')
    if not refs:
        return
    _section_title(doc, "References & Further Study", "\U0001F4DA")
    intro = ("All Spirit of Prophecy works below are freely available to read, "
             "search, and verify at the Ellen G. White Estate website "
             "(egwwritings.org). Scripture quotations are from the King James "
             "Version (KJV) unless noted; the full text is available at "
             "biblegateway.com.")
    render_paragraph(doc, intro)
    for r in refs:
        if isinstance(r, (tuple, list)):
            citation, url = (r + ("",))[:2]
        else:
            citation, url = r.get('citation', ''), r.get('url', '')
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.32)
        pf.first_line_indent = Inches(-0.22)
        pf.space_after = Pt(4)
        run = p.add_run("•  "); run.font.color.rgb = RGBColor.from_string(GREEN_SOFT)
        add_runs(p, citation, color=INK, size=10.5)
        if url:
            add_runs(p, f"  {url}", color=GREEN_SOFT, size=9.5)


# ────────────────────────────── public API ────────────────────────────────
def build_lesson(meta, out_path):
    """Render one lesson dictionary to a .docx file at *out_path*."""
    doc = Document()
    _configure_styles(doc)
    _set_margins(doc)
    _header_footer(doc, meta)

    _title_block(doc, meta)
    _glance_panel(doc, meta)
    _memory_verse(doc, meta)

    if meta.get('intro'):
        for blk in meta['intro']:
            render_block(doc, blk)

    for blk in meta.get('body', []):
        render_block(doc, blk)

    if meta.get('summary'):
        _section_title(doc, "Lesson Summary", "\U0001F4DD")
        render_callout(doc, None, [("ul", meta['summary'])])

    if meta.get('review_questions'):
        _section_title(doc, "Review & Reflection", "❓")
        render_list(doc, meta['review_questions'], ordered=True)

    if meta.get('activity'):
        _section_title(doc, "Practical Assignment", "\U0001F6E0")
        for blk in meta['activity']:
            render_block(doc, blk)

    _references(doc, meta)

    doc.save(out_path)
    return out_path
